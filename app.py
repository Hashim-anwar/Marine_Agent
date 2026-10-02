"""MarineWise AI - simple Streamlit MVP for marine engine troubleshooting and training."""

from __future__ import annotations

import io
import os
import re
from typing import Any

import streamlit as st
from groq import Groq
from PIL import Image

from agents import make_assessment_agent, make_training_agent, make_troubleshooting_agent
from rag import (
    DEFAULT_CHUNK_SIZE,
    DEFAULT_CHUNK_OVERLAP,
    build_index,
    download_google_drive_pdf,
    get_embedder,
    retrieve_context,
    search_index,
)

MODEL = "openai/gpt-oss-120b"

st.set_page_config(page_title="MarineWise AI", page_icon="⚓", layout="wide")

st.markdown(
    """
    <style>
    .stApp { background: #f5f8fb; }
    [data-testid="stSidebar"] { background: #0b2239; }
    [data-testid="stSidebar"] * { color: #eef6ff; }
    h1, h2, h3 { color: #0b3558; }
    .marine-card { padding: 1rem 1.2rem; border-radius: 12px; background: white; border: 1px solid #dbe5ee; margin-bottom: 1rem; }
    .source-tag { padding: .25rem .5rem; border-radius: 6px; background: #e8f2fb; color: #0b4a75; display: inline-block; margin: .1rem; font-size: .85rem; }
    </style>
    """,
    unsafe_allow_html=True,
)


def get_api_key() -> str | None:
    """Read the Groq key only from Streamlit secrets or the environment."""
    try:
        key = st.secrets.get("GROQ_API_KEY")
    except Exception:
        key = None
    return key or os.getenv("GROQ_API_KEY")


def get_client() -> Groq:
    key = get_api_key()
    if not key:
        raise RuntimeError("GROQ_API_KEY is not configured. Add it to Streamlit secrets or your environment.")
    return Groq(api_key=key)


@st.cache_resource(show_spinner=False)
def cached_embedder():
    return get_embedder()


def ask_groq(system_prompt: str, user_prompt: str, browser_search: bool = False) -> str:
    client = get_client()
    kwargs: dict[str, Any] = {
        "model": MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.2,
        "max_tokens": 3000,
        "reasoning_effort": "medium",
    }
    if browser_search:
        kwargs["tools"] = [{"type": "browser_search"}]
    response = client.chat.completions.create(**kwargs)
    return response.choices[0].message.content or "No answer was returned."


def make_context_query(manufacturer: str, model: str, alarm: str) -> str:
    return f"Marine engine manufacturer {manufacturer}; engine model {model}; defect/alarm: {alarm}"


def render_sources(results: list[dict[str, Any]]) -> None:
    if not results:
        return
    st.markdown("**Manual sources retrieved:**")
    unique = []
    seen = set()
    for r in results:
        key = (r["source"], r["page"])
        if key not in seen:
            seen.add(key)
            unique.append(key)
    for source, page in unique:
        st.markdown(f'<span class="source-tag">{source} — page {page}</span>', unsafe_allow_html=True)


def sidebar_manuals() -> None:
    st.sidebar.markdown("## ⚓ MarineWise AI")
    st.sidebar.caption("Marine engine troubleshooting + technician training")
    st.sidebar.markdown("### Manuals")
    uploads = st.sidebar.file_uploader(
        "Upload PDF manuals",
        type=["pdf"],
        accept_multiple_files=True,
        help="Upload one or more manuals. The app only indexes them when you click Build / Rebuild Index.",
    )
    drive_url = st.sidebar.text_input("Or paste a public Google Drive PDF link", placeholder="https://drive.google.com/file/d/...")
    if st.sidebar.button("Build / Rebuild FAISS Index", type="primary", use_container_width=True):
        items: list[tuple[str, bytes]] = []
        if uploads:
            items.extend((f.name, f.getvalue()) for f in uploads)
        if drive_url.strip():
            try:
                with st.spinner("Downloading the Google Drive PDF..."):
                    items.append(download_google_drive_pdf(drive_url.strip()))
            except Exception as exc:
                st.sidebar.error(f"Drive download failed: {exc}")
        if not items:
            st.sidebar.warning("Upload a PDF or provide a public Google Drive PDF link first.")
        else:
            try:
                with st.spinner("Reading manuals, creating chunks, and building FAISS..."):
                    st.session_state.rag = build_index(items, cached_embedder())
                st.sidebar.success(f"Indexed {len(items)} manual(s).")
            except Exception as exc:
                st.sidebar.error(f"Indexing failed: {exc}")

    rag = st.session_state.get("rag")
    if rag:
        st.sidebar.success(f"Index ready: {len(rag['records'])} chunks")
        st.sidebar.caption(f"Chunk size: {rag['chunk_size']} characters")
    else:
        st.sidebar.info("No manual index yet. Upload/paste a source and build the index when needed.")


def troubleshooting_page() -> None:
    st.title("Troubleshooting Agent")
    st.write("Search your supplied manuals first. The answer is deliberately grounded in the retrieved manual pages.")

    with st.form("troubleshoot_form"):
        manufacturer = st.selectbox("Manufacturer", ["CAT", "MTU", "YANMAR", "YAMAHA", "HONDA", "MAN", "OTHER"])
        engine_model = st.text_input("Engine Model", placeholder="e.g. MTU 16V 4000 M90")
        serial = st.text_input("Serial Number (optional)")
        defect = st.text_area("Defect or Alarm", placeholder="Describe the symptom, alarm code, and what happened.")
        submitted = st.form_submit_button("Troubleshoot", type="primary")

    if not submitted:
        return
    if not engine_model.strip() or not defect.strip():
        st.warning("Please enter the engine model and defect/alarm.")
        return

    rag = st.session_state.get("rag")
    if not rag:
        st.warning("No manuals are indexed yet. Use the sidebar to upload manuals and build the FAISS index.")
        return

    query = make_context_query(manufacturer, engine_model, defect)
    with st.spinner("Searching the manuals..."):
        results = search_index(rag, query, cached_embedder(), k=6)
        context, relevant = retrieve_context(results)
        st.session_state.last_retrieved = relevant

    st.metric("Retrieved manual chunks", len(relevant))
    st.caption(f"Chunk size: {rag['chunk_size']} characters • Retrieved pages: {len({r['page'] for r in relevant})}")
    render_sources(relevant)

    if not context:
        st.warning("Not found in manuals. Do you want me to search online?")
        if st.button("Yes — search online", key="web_troubleshoot"):
            with st.spinner("Searching online with Groq browser search..."):
                answer = ask_groq(
                    "You are a marine engine troubleshooting assistant. Search the web for reliable manufacturer documentation or reputable technical sources. Do not invent specifications. Clearly state that the answer is from the web, give links/source names when available, and advise the user to verify against the engine's current manual.",
                    f"Find reliable information for {manufacturer} {engine_model}, serial {serial or 'not provided'}, symptom/alarm: {defect}. Explain likely checks and safe next steps.",
                    browser_search=True,
                )
            st.markdown("### Web-sourced answer")
            st.write(answer)
        return

    with st.spinner("Preparing a manual-grounded answer..."):
        answer = ask_groq(
            "You are a marine engine troubleshooting assistant. Use ONLY the supplied manual excerpts. Do not invent facts, specifications, causes, limits, or procedures. If the excerpts do not answer the question, say exactly: 'Not found in manuals. Do you want me to search online?' Cite the source file and exact page number for claims.",
            f"Manufacturer: {manufacturer}\nEngine model: {engine_model}\nSerial: {serial or 'not provided'}\nDefect/alarm: {defect}\n\nMANUAL EXCERPTS:\n{context}",
        )
    st.markdown("### Troubleshooting answer")
    st.write(answer)


def training_material_page() -> None:
    st.subheader("2A. Training Material")
    with st.form("training_material_form"):
        engine = st.text_input("Engine Model", placeholder="MTU 16V 4000 M90")
        ship = st.text_input("Ship", placeholder="MV Example")
        topic = st.text_input("Training Topic", placeholder="Fuel Injection System")
        c1, c2, c3 = st.columns(3)
        make_pdf = c1.checkbox("PDF", True)
        make_ppt = c2.checkbox("PPT", True)
        make_word = c3.checkbox("Word", False)
        submitted = st.form_submit_button("Generate Training Material", type="primary")

    if not submitted:
        return
    if not engine.strip() or not topic.strip():
        st.warning("Enter an engine model and training topic.")
        return
    if not any([make_pdf, make_ppt, make_word]):
        st.warning("Select at least one output format.")
        return

    rag = st.session_state.get("rag")
    context = ""
    relevant: list[dict[str, Any]] = []
    if rag:
        with st.spinner("Searching manuals first..."):
            results = search_index(rag, f"{engine} {topic}", cached_embedder(), k=8)
            context, relevant = retrieve_context(results, min_score=0.25)
            st.session_state.last_retrieved = relevant
    else:
        st.info("No manual index is available, so this will be general training content. For manual-grounded training, build the index first.")

    with st.spinner("Generating training content..."):
        content = ask_groq(
            "You are a marine technical training instructor. Create practical, beginner-friendly technician training material. Prefer the supplied manual excerpts when present. Never invent manual-specific values. Clearly label general knowledge when the manuals do not cover something. Include learning objectives, system overview, operating principle, inspection/check steps, common faults, safety notes, and a short knowledge check. Use concise headings and bullets.",
            f"Engine: {engine}\nShip: {ship or 'Not specified'}\nTopic: {topic}\n\nMANUAL EXCERPTS (may be empty):\n{context}",
        )

    st.markdown("### Training material")
    st.write(content)
    render_sources(relevant)
    st.caption(f"Chunk size: {rag['chunk_size'] if rag else DEFAULT_CHUNK_SIZE} • Retrieved pages: {len({r['page'] for r in relevant})}")

    diagram = make_training_diagram(engine, topic)
    st.image(diagram, caption="Simple system-learning diagram generated for the training package.")

    if make_pdf:
        st.download_button("Download PDF", make_training_pdf(engine, ship, topic, content, relevant, diagram), "marinewise_training.pdf", "application/pdf")
    if make_ppt:
        st.download_button("Download PowerPoint", make_training_ppt(engine, ship, topic, content, diagram), "marinewise_training.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")
    if make_word:
        st.download_button("Download Word", make_training_docx(engine, ship, topic, content, diagram), "marinewise_training.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def quiz_page() -> None:
    st.subheader("2B. Quiz Generator")
    with st.form("quiz_form"):
        topic = st.text_input("Topic", placeholder="Fuel Injection System")
        qtype = st.selectbox("Type", ["MCQ", "Short Question", "True-False"])
        count = st.number_input("Number of questions", min_value=1, max_value=30, value=10, step=1)
        submitted = st.form_submit_button("Generate Quiz", type="primary")
    if not submitted:
        return
    if not topic.strip():
        st.warning("Enter a topic.")
        return

    rag = st.session_state.get("rag")
    context = ""
    relevant: list[dict[str, Any]] = []
    if rag:
        results = search_index(rag, topic, cached_embedder(), k=8)
        context, relevant = retrieve_context(results, min_score=0.25)
        st.session_state.last_retrieved = relevant

    with st.spinner("Generating quiz and answer key..."):
        quiz_text = ask_groq(
            "Create a technician training quiz. Follow the requested type and count exactly. Include an ANSWER KEY at the end. Base it on supplied manual excerpts where available and do not invent manual-specific values.",
            f"Topic: {topic}\nType: {qtype}\nQuestions: {count}\n\nMANUAL EXCERPTS:\n{context}",
        )
    st.markdown("### Quiz")
    st.write(quiz_text)
    st.download_button("Download Quiz PDF", make_quiz_pdf(topic, qtype, quiz_text), "marinewise_quiz.pdf", "application/pdf")
    render_sources(relevant)


def assessment_page() -> None:
    st.subheader("2C. Score an Assessment")
    st.write("Upload a clear photo or scan. OCR quality affects scoring, so review the extracted text before relying on the score.")
    uploaded = st.file_uploader("Assessment image", type=["jpg", "jpeg", "png"])
    answer_key = st.text_area("Optional answer key", placeholder="Example: 1=A, 2=C, 3=True, 4=B")
    if uploaded and st.button("Score Assessment", type="primary"):
        with st.spinner("Reading the assessment image..."):
            try:
                import pytesseract
                image = Image.open(uploaded).convert("RGB")
                st.image(image, caption="Uploaded assessment", width=500)
                extracted = pytesseract.image_to_string(image)
            except Exception as exc:
                st.error(f"OCR could not run: {exc}. Install Tesseract OCR locally or use a clearer image/environment with OCR support.")
                return

        st.text_area("OCR text — check this before scoring", extracted, height=250)
        if not answer_key.strip():
            st.warning("Add an answer key so the app can calculate a defensible score.")
            return
        with st.spinner("Scoring the assessment..."):
            score = score_with_llm(extracted, answer_key)
        st.markdown("### Assessment result")
        st.write(score["feedback"])
        st.metric("Score", f"{score['score']:.0f}%")
        if score["score"] < 50:
            st.warning("Below 50% — remedial training is recommended.")
            with st.spinner("Generating remedial presentation..."):
                remedial = ask_groq(
                    "Create a short remedial marine technician training presentation outline based only on the assessment mistakes. Include 5-7 slides, explanations, practice checks, and a final retest. Do not invent technical values.",
                    f"Assessment OCR:\n{extracted}\n\nAnswer key:\n{answer_key}\n\nScoring feedback:\n{score['feedback']}",
                )
            st.download_button("Download Remedial Training PPT", make_remedial_ppt(remedial), "marinewise_remedial_training.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation")


def score_with_llm(extracted: str, answer_key: str) -> dict[str, Any]:
    raw = ask_groq(
        "Score a technician assessment. Parse the provided answer key. Return exactly three lines: SCORE_PERCENT: number; FEEDBACK: concise explanation; MISSED_TOPICS: comma-separated topics. Do not guess unreadable answers; mark them incorrect/unclear and explain.",
        f"OCR answers:\n{extracted}\n\nAnswer key:\n{answer_key}",
    )
    match = re.search(r"SCORE_PERCENT:\s*([0-9]+(?:\.[0-9]+)?)", raw, re.I)
    score = float(match.group(1)) if match else 0.0
    feedback = raw.split("FEEDBACK:", 1)[1].split("MISSED_TOPICS:", 1)[0].strip() if "FEEDBACK:" in raw else raw
    return {"score": max(0.0, min(100.0, score)), "feedback": feedback}


def learning_page() -> None:
    st.title("Learning: RAG Basics")
    st.markdown("### RAG")
    st.write("Retrieval-Augmented Generation (RAG) first searches your documents, then gives the most relevant passages to the language model. This reduces unsupported answers.")
    st.markdown("### Advanced RAG")
    st.write("Advanced RAG can add query rewriting, metadata filters, reranking, multiple retrieval steps, or better chunking. This MVP keeps those ideas simple.")
    st.markdown("### Chunks")
    st.write("A chunk is a small piece of a manual. This MVP uses about 900 characters with 120 characters of overlap so nearby sentences are not separated too aggressively.")
    st.markdown("### FAISS")
    st.write("FAISS is a fast vector-search library. Each chunk is converted into a numeric embedding; FAISS finds chunks whose vectors are closest to the question.")
    st.code("Manual PDF → pages → chunks → embeddings → FAISS\n                                  ↓\nQuestion → embedding → nearest chunks → Groq → answer + page/source")
    rag = st.session_state.get("rag")
    c1, c2, c3 = st.columns(3)
    c1.metric("Chunk size", rag["chunk_size"] if rag else DEFAULT_CHUNK_SIZE)
    c2.metric("Number of chunks", len(rag["records"]) if rag else 0)
    c3.metric("Retrieved pages", len({r["page"] for r in st.session_state.get("last_retrieved", [])}))
    if rag:
        st.caption(f"Embedding model: {rag['embedding_model']} • Overlap: {rag['chunk_overlap']} characters")


def make_training_diagram(engine: str, topic: str) -> bytes:
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch

    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")
    boxes = [(0.5, 1.8, "Fuel / Air\nInput"), (3.0, 1.8, "Engine\nCombustion"), (5.5, 1.8, "Sensors /\nControls"), (8.0, 1.8, "Alarm /\nOutput")]
    for x, y, label in boxes:
        ax.add_patch(FancyBboxPatch((x, y), 1.7, 1.2, boxstyle="round,pad=0.03"))
        ax.text(x + 0.85, y + 0.6, label, ha="center", va="center", fontsize=10)
    for x in [2.2, 4.7, 7.2]:
        ax.annotate("", xy=(x + 0.6, 2.4), xytext=(x, 2.4), arrowprops={"arrowstyle": "->", "lw": 1.5})
    ax.text(5, 4.2, f"MarineWise training: {engine} — {topic}", ha="center", fontsize=13, fontweight="bold")
    buffer = io.BytesIO()
    fig.savefig(buffer, format="png", dpi=160, bbox_inches="tight")
    plt.close(fig)
    return buffer.getvalue()


def split_text(text: str, max_chars: int = 1800) -> list[str]:
    parts = []
    current = ""
    for para in text.split("\n"):
        para = para.strip()
        if not para:
            continue
        if len(current) + len(para) + 1 > max_chars:
            if current:
                parts.append(current)
            current = para
        else:
            current += ("\n" if current else "") + para
    if current:
        parts.append(current)
    return parts or ["No content generated."]


def make_training_pdf(engine: str, ship: str, topic: str, content: str, sources: list[dict[str, Any]], diagram: bytes) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Image as RLImage, Paragraph, SimpleDocTemplate, Spacer

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = [Paragraph("MarineWise AI — Technical Training", styles["Title"]), Paragraph(f"Engine: {engine} | Ship: {ship or 'Not specified'} | Topic: {topic}", styles["Heading2"]), Spacer(1, 10), RLImage(io.BytesIO(diagram), width=6.8 * inch, height=3.06 * inch)]
    for part in split_text(content):
        story.extend([Paragraph(part.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)])
    if sources:
        story.append(Paragraph("Manual sources", styles["Heading2"]))
        for s in sources:
            story.append(Paragraph(f"{s['source']} — page {s['page']}", styles["BodyText"]))
    doc.build(story)
    return out.getvalue()


def make_training_docx(engine: str, ship: str, topic: str, content: str, diagram: bytes) -> bytes:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("MarineWise AI — Technical Training", 0)
    doc.add_paragraph(f"Engine: {engine} | Ship: {ship or 'Not specified'} | Topic: {topic}")
    doc.add_picture(io.BytesIO(diagram), width=Inches(6.5))
    for part in split_text(content):
        doc.add_paragraph(part)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def make_training_ppt(engine: str, ship: str, topic: str, content: str, diagram: bytes) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "MarineWise AI — Technical Training"
    title.placeholders[1].text = f"{engine} | {ship or 'Not specified'} | {topic}"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "System Overview"
    slide.shapes.add_picture(io.BytesIO(diagram), Inches(0.7), Inches(1.4), width=Inches(8.6))
    for i, part in enumerate(split_text(content, 1100)[:8], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Training — Part {i}"
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(8.6), Inches(5.5))
        tf = box.text_frame
        tf.text = part
        for p in tf.paragraphs:
            p.font.size = Pt(18)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def make_quiz_pdf(topic: str, qtype: str, text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    story = [Paragraph("MarineWise AI — Technician Quiz", styles["Title"]), Paragraph(f"Topic: {topic} | Type: {qtype}", styles["Heading2"]), Spacer(1, 10)]
    for part in split_text(text, 1500):
        story.extend([Paragraph(part.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>"), styles["BodyText"]), Spacer(1, 8)])
    doc.build(story)
    return out.getvalue()


def make_remedial_ppt(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    for i, part in enumerate(split_text(text, 1000)[:8], start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = f"Remedial Training — Slide {i}"
        box = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(8.6), Inches(5.5))
        box.text_frame.text = part
        for p in box.text_frame.paragraphs:
            p.font.size = Pt(18)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def main() -> None:
    sidebar_manuals()
    page = st.sidebar.radio("Navigate", ["Troubleshooting Agent", "Technical Training", "Learning"])
    if page == "Troubleshooting Agent":
        troubleshooting_page()
    elif page == "Technical Training":
        tabs = st.tabs(["2A Training Material", "2B Quiz Generator", "2C Score Assessment"])
        with tabs[0]:
            training_material_page()
        with tabs[1]:
            quiz_page()
        with tabs[2]:
            assessment_page()
    else:
        learning_page()


if __name__ == "__main__":
    main()
